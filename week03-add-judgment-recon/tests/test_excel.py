"""The workbook. Checked by reading the file back, not by trusting the writer."""

from __future__ import annotations

import csv
from datetime import datetime
from decimal import Decimal

import openpyxl
import pytest

from recon.engine import reconcile
from recon.excel import MONEY_FORMAT, OutputLocked, write_workbook
from recon.model import Side

from .helpers import make_file


@pytest.fixture
def workbook(gl, bank, tmp_path):
    result = reconcile(gl, bank)
    path = write_workbook(result, tmp_path / "recon.xlsx")
    return openpyxl.load_workbook(path), result


def summary_pairs(wb) -> dict:
    """Column A to column B, for the label/value rows on the Summary tab."""
    out = {}
    for label, value in wb["Summary"].iter_rows(
        min_row=1, max_row=wb["Summary"].max_row, max_col=2, values_only=True
    ):
        if label is not None and value is not None:
            out[str(label).strip()] = value
    return out


# --- structure -----------------------------------------------------------

def test_expected_sheets_in_working_order(workbook):
    wb, _ = workbook
    assert wb.sheetnames == [
        "Summary", "Reconciled", "Exceptions", "General Ledger", "Bank",
    ]


def test_no_ambiguous_tab_when_nothing_was_declined(workbook):
    """Absence means there was nothing to decline, not nothing to say."""
    wb, result = workbook
    assert result.ambiguities == []
    assert "Ambiguous" not in wb.sheetnames


def test_ambiguous_tab_appears_when_a_pass_declines(tmp_path):
    g = make_file(Side.GL, [(5, 5000, "Inv 700"), (5, 5000, "Inv 700")])
    b = make_file(Side.BANK, [(5, 5000, "Inv 700")])
    result = reconcile(g, b)
    assert result.ambiguities

    wb = openpyxl.load_workbook(write_workbook(result, tmp_path / "amb.xlsx"))
    assert "Ambiguous" in wb.sheetnames
    ws = wb["Ambiguous"]
    assert [c.value for c in ws[1]] == [
        "pass", "key", "gl_rows", "bank_rows", "why not claimed",
        "likely cause", "suggested check", "explanation", "explained by",
    ]
    assert ws.cell(row=2, column=3).value == "GL-0001, GL-0002"
    assert "indistinguishable" in ws.cell(row=2, column=5).value
    # No explanations were passed, so the advisory columns stay empty.
    assert all(ws.cell(row=2, column=c).value is None for c in (6, 7, 8, 9))


def test_tables_are_filterable_and_frozen(workbook):
    wb, _ = workbook
    for name in ("Reconciled", "Exceptions", "General Ledger", "Bank"):
        ws = wb[name]
        assert ws.auto_filter.ref is not None, name
        assert ws.freeze_panes == "A2", name


# --- the source tabs are the source --------------------------------------

def test_source_tabs_keep_the_original_columns_first(workbook):
    wb, _ = workbook
    assert [c.value for c in wb["General Ledger"][1]] == [
        "gl_date", "amount", "main_memo", "line_memo",
        "row_id", "match_id", "match_shape", "match_rule",
    ]
    assert [c.value for c in wb["Bank"][1]] == [
        "posted_date", "posted_amount", "messy_memo",
        "row_id", "match_id", "match_shape", "match_rule",
    ]


def test_source_values_are_verbatim(workbook, corpus_dir):
    """Byte-for-byte what the CSV held, including the amount as written."""
    wb, _ = workbook
    with (corpus_dir / "general_ledger.csv").open(encoding="utf-8-sig") as fh:
        original = list(csv.DictReader(fh))

    ws = wb["General Ledger"]
    assert ws.max_row == len(original) + 1
    for i, src in enumerate(original, start=2):
        assert ws.cell(row=i, column=1).value == src["gl_date"]
        assert ws.cell(row=i, column=2).value == src["amount"]
        assert ws.cell(row=i, column=3).value == src["main_memo"]
        assert ws.cell(row=i, column=4).value == src["line_memo"]


def test_source_rows_carry_their_match(workbook):
    wb, result = workbook
    matched = {}
    for m in result.matches:
        for row_id in m.gl_ids:
            matched[row_id] = m.match_id

    ws = wb["General Ledger"]
    unmatched_seen = 0
    for i in range(2, ws.max_row + 1):
        row_id = ws.cell(row=i, column=5).value
        written = ws.cell(row=i, column=6).value
        if row_id in matched:
            assert written == matched[row_id]
        else:
            # An open item leaves the cell blank, which is what Excel's own
            # "Blanks" filter looks for. openpyxl reads that back as None.
            assert written is None
            unmatched_seen += 1
    assert unmatched_seen == len(result.unmatched_gl)


# --- every row is accounted for exactly once -----------------------------

def test_all_rows_appear_once_across_reconciled_and_exceptions(workbook):
    wb, result = workbook
    ids = [r[4] for r in wb["Reconciled"].iter_rows(min_row=2, values_only=True)]
    ids += [r[1] for r in wb["Exceptions"].iter_rows(min_row=2, values_only=True)]
    expected = len(result.gl.rows) + len(result.bank.rows)
    assert len(ids) == expected
    assert len(set(ids)) == expected


def test_group_members_share_one_match_id(workbook):
    """A five-to-one group is five rows plus one, all under one id."""
    wb, _ = workbook
    rows = list(wb["Reconciled"].iter_rows(min_row=2, values_only=True))
    groups: dict[str, list[tuple]] = {}
    for r in rows:
        groups.setdefault(r[0], []).append(r)

    five_to_one = [g for g in groups.values() if g[0][1] == "5:1"]
    assert len(five_to_one) == 5
    for group in five_to_one:
        assert len(group) == 6
        assert sum(1 for r in group if r[3] == "GL") == 5
        assert sum(1 for r in group if r[3] == "Bank") == 1
        # The five ledger rows sum to the one deposit.
        gl_total = sum(Decimal(str(r[7])) for r in group if r[3] == "GL")
        bank_total = sum(Decimal(str(r[7])) for r in group if r[3] == "Bank")
        assert gl_total == bank_total


def test_exceptions_are_the_open_items(workbook):
    wb, result = workbook
    rows = list(wb["Exceptions"].iter_rows(min_row=2, values_only=True))
    assert len(rows) == len(result.unmatched_gl) + len(result.unmatched_bank)
    assert {r[0] for r in rows} == {"GL", "Bank"}
    assert all("open item" in r[7] for r in rows)


# --- the proof, recomputed from the file --------------------------------

def test_workbook_foots_on_its_own_numbers(workbook, corpus_dir):
    """Add up what the workbook says and check it against the source CSVs.

    Deliberately independent of the engine: if the writer dropped or duplicated
    a row, these totals would not tie even though reconcile() was happy.
    """
    wb, _ = workbook
    rec = list(wb["Reconciled"].iter_rows(min_row=2, values_only=True))
    exc = list(wb["Exceptions"].iter_rows(min_row=2, values_only=True))

    gl_matched = sum(Decimal(str(r[7])) for r in rec if r[3] == "GL")
    bank_matched = sum(Decimal(str(r[7])) for r in rec if r[3] == "Bank")
    gl_open = sum(Decimal(str(r[4])) for r in exc if r[0] == "GL")
    bank_open = sum(Decimal(str(r[4])) for r in exc if r[0] == "Bank")

    with (corpus_dir / "general_ledger.csv").open(encoding="utf-8-sig") as fh:
        gl_total = sum(Decimal(r["amount"]) for r in csv.DictReader(fh))
    with (corpus_dir / "bank.csv").open(encoding="utf-8-sig") as fh:
        bank_total = sum(Decimal(r["posted_amount"]) for r in csv.DictReader(fh))

    assert gl_matched + gl_open == gl_total
    assert bank_matched + bank_open == bank_total

    drift = gl_matched - bank_matched
    assert (gl_open - bank_open) + drift == gl_total - bank_total


def test_summary_reports_the_verdict(workbook):
    wb, _ = workbook
    pairs = summary_pairs(wb)
    assert pairs["Result"].startswith("FOOTS")
    assert pairs["Difference to explain"] == pytest.approx(1027.21)
    assert pairs["Drift accepted inside matches"] == pytest.approx(-0.11)
    assert pairs["Accounted for"] == pytest.approx(1027.21)
    assert pairs["GL total less bank total"] == pytest.approx(1027.21)


def test_summary_accounts_for_every_pass(workbook):
    wb, result = workbook
    pairs = summary_pairs(wb)
    for name, count in result.pass_counts:
        assert pairs[name] == count
    assert pairs["Match groups"] == len(result.matches)
    assert pairs["Ambiguous candidates"] == 0


# --- cell types ----------------------------------------------------------

def test_amounts_are_numbers_not_text(workbook):
    """A reconciliation you cannot sum in Excel is not much use."""
    wb, _ = workbook
    for name, col in (("Reconciled", 8), ("Exceptions", 5)):
        ws = wb[name]
        for i in range(2, min(ws.max_row, 12) + 1):
            cell = ws.cell(row=i, column=col)
            assert isinstance(cell.value, (int, float)), (name, i, cell.value)
            assert cell.number_format == MONEY_FORMAT


def test_dates_are_dates_not_text(workbook):
    wb, _ = workbook
    for name, col in (("Reconciled", 7), ("Exceptions", 4)):
        ws = wb[name]
        for i in range(2, min(ws.max_row, 12) + 1):
            assert isinstance(ws.cell(row=i, column=col).value, datetime)


def test_negative_amounts_survive_the_round_trip(workbook):
    """The void legs must not be lost or absolutised on the way out."""
    wb, _ = workbook
    amounts = [
        r[7] for r in wb["Reconciled"].iter_rows(min_row=2, values_only=True)
    ]
    assert -100.0 in amounts   # set 9, the GL void
    assert -300.0 in amounts   # set 10, the bank void


# --- failure handling ----------------------------------------------------

def test_locked_output_is_a_clear_message_not_a_traceback(gl, bank, tmp_path, monkeypatch):
    """The commonest failure by far: the workbook is already open in Excel."""
    from openpyxl.workbook import workbook as wb_module

    def refuse(self, filename):
        raise PermissionError(13, "Permission denied")

    monkeypatch.setattr(wb_module.Workbook, "save", refuse)
    with pytest.raises(OutputLocked, match="probably open in Excel"):
        write_workbook(reconcile(gl, bank), tmp_path / "locked.xlsx")


def test_creates_missing_directories(gl, bank, tmp_path):
    target = tmp_path / "nested" / "deeper" / "recon.xlsx"
    assert write_workbook(reconcile(gl, bank), target) == target
    assert target.is_file()
