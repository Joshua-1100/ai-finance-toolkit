"""The Excel workbook.

Five tabs, in the order you would work them:

  Summary          the control totals and the proof that the recon is complete
  Reconciled       every matched row, grouped, with the quality of its match
  Exceptions       what is still open - the list a Controller actually works
  General Ledger   the source file, verbatim, plus the match it landed in
  Bank             the same for the statement

An Ambiguous tab appears only when a pass declined to choose between candidates.
Its absence means there was nothing to decline.

Amounts are written as numbers, not text, because a reconciliation you cannot sum
in Excel is not much use. That means Excel holds them as float64 - which is fine
for display and totalling at two decimals, and is not where the matching
happened. Every comparison that decided anything was made in integer cents
before this file existed.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side as BorderSide
from openpyxl.utils import get_column_letter

from .engine import Result
from .model import LedgerRow, Side

MONEY_FORMAT = "#,##0.00;[Red]-#,##0.00"
DATE_FORMAT = "yyyy-mm-dd"

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", start_color="44546A")
TITLE_FONT = Font(bold=True, size=12)
SECTION_FONT = Font(bold=True)
BAND_FILL = PatternFill("solid", start_color="F2F5F9")
FOOTS_FONT = Font(bold=True, color="006100")
FAILS_FONT = Font(bold=True, color="9C0006")
TOP_BORDER = Border(top=BorderSide(style="thin"))

MAX_WIDTH = 46


class OutputLocked(Exception):
    """The workbook could not be written - almost always open in Excel."""


def _write_header(ws, headers: list[str]) -> None:
    for col, name in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"


def _finish_table(ws, headers: list[str], widths: list[int]) -> None:
    """Autofilter and column widths, once the rows are in."""
    last_col = get_column_letter(len(headers))
    ws.auto_filter.ref = f"A1:{last_col}{max(ws.max_row, 1)}"
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = min(width, MAX_WIDTH)


def _money(ws, row: int, col: int, cents: int):
    cell = ws.cell(row=row, column=col, value=cents / 100)
    cell.number_format = MONEY_FORMAT
    return cell


def _date(ws, row: int, col: int, value):
    cell = ws.cell(row=row, column=col, value=value)
    cell.number_format = DATE_FORMAT
    return cell


# --- Reconciled ----------------------------------------------------------

RECONCILED_HEADERS = [
    "match_id", "shape", "rule", "side", "row_id", "source_row", "date",
    "amount", "memo", "main_memo", "amount_match", "date_match",
    "identity_match", "note",
]
RECONCILED_WIDTHS = [10, 7, 17, 6, 9, 10, 12, 13, 26, 20, 13, 11, 14, 46]


def _write_reconciled(ws, result: Result) -> None:
    _write_header(ws, RECONCILED_HEADERS)
    rows_by_id = {
        r.row_id: r for r in list(result.gl.rows) + list(result.bank.rows)
    }

    excel_row = 2
    for group_index, match in enumerate(result.matches):
        # Alternate a fill per group, so a five-row group reads as one block
        # rather than five unrelated lines.
        banded = group_index % 2 == 1
        members = [(Side.GL, i) for i in match.gl_ids]
        members += [(Side.BANK, i) for i in match.bank_ids]

        for side, row_id in members:
            row: LedgerRow = rows_by_id[row_id]
            ws.cell(row=excel_row, column=1, value=match.match_id)
            ws.cell(row=excel_row, column=2, value=match.shape)
            ws.cell(row=excel_row, column=3, value=match.rule)
            ws.cell(row=excel_row, column=4,
                    value="GL" if side is Side.GL else "Bank")
            ws.cell(row=excel_row, column=5, value=row.row_id)
            ws.cell(row=excel_row, column=6, value=row.source_row)
            _date(ws, excel_row, 7, row.txn_date)
            _money(ws, excel_row, 8, row.amount_cents)
            ws.cell(row=excel_row, column=9, value=row.memo)
            ws.cell(row=excel_row, column=10, value=row.main_memo)
            ws.cell(row=excel_row, column=11, value=match.amount.value)
            ws.cell(row=excel_row, column=12, value=match.date.value)
            ws.cell(row=excel_row, column=13, value=match.identity.value)
            ws.cell(row=excel_row, column=14, value=match.note)

            if banded:
                for col in range(1, len(RECONCILED_HEADERS) + 1):
                    ws.cell(row=excel_row, column=col).fill = BAND_FILL
            excel_row += 1

    _finish_table(ws, RECONCILED_HEADERS, RECONCILED_WIDTHS)


# --- Exceptions ----------------------------------------------------------

EXCEPTION_HEADERS = [
    "side", "row_id", "source_row", "date", "amount", "memo", "main_memo",
    "reason",
]
EXCEPTION_WIDTHS = [6, 9, 10, 12, 13, 26, 20, 40]

NO_MATCH = "No match found by any pass - open item"


def _write_exceptions(ws, result: Result) -> None:
    _write_header(ws, EXCEPTION_HEADERS)
    excel_row = 2
    for side, rows in (("GL", result.unmatched_gl), ("Bank", result.unmatched_bank)):
        for row in rows:
            ws.cell(row=excel_row, column=1, value=side)
            ws.cell(row=excel_row, column=2, value=row.row_id)
            ws.cell(row=excel_row, column=3, value=row.source_row)
            _date(ws, excel_row, 4, row.txn_date)
            _money(ws, excel_row, 5, row.amount_cents)
            ws.cell(row=excel_row, column=6, value=row.memo)
            ws.cell(row=excel_row, column=7, value=row.main_memo)
            ws.cell(row=excel_row, column=8, value=NO_MATCH)
            excel_row += 1
    _finish_table(ws, EXCEPTION_HEADERS, EXCEPTION_WIDTHS)


# --- Ambiguous -----------------------------------------------------------

AMBIGUOUS_HEADERS = [
    "pass", "key", "gl_rows", "bank_rows", "why not claimed",
    "likely cause", "suggested check", "explanation", "explained by",
]
AMBIGUOUS_WIDTHS = [18, 28, 24, 24, 44, 20, 24, 46, 18]


def _write_ambiguous(ws, result: Result, explanations: dict | None = None) -> None:
    """Candidates the engine declined, with the model's read on each.

    The last four columns are advisory. Nothing in them affected a match - the
    engine had already finished and proved itself before they were written.
    """
    _write_header(ws, AMBIGUOUS_HEADERS)
    explanations = explanations or {}

    for excel_row, amb in enumerate(result.ambiguities, start=2):
        ws.cell(row=excel_row, column=1, value=amb.pass_name)
        ws.cell(row=excel_row, column=2, value=amb.key)
        ws.cell(row=excel_row, column=3, value=", ".join(amb.gl_ids))
        ws.cell(row=excel_row, column=4, value=", ".join(amb.bank_ids))
        ws.cell(row=excel_row, column=5, value=amb.reason)

        found = explanations.get(f"{amb.pass_name}|{amb.key}")
        if found is None:
            continue
        ws.cell(row=excel_row, column=6, value=found.cause_label)
        ws.cell(row=excel_row, column=7, value=found.check_label)
        ws.cell(row=excel_row, column=8, value=found.summary)
        ws.cell(row=excel_row, column=9,
                value="not available" if found.failed else found.model)

    _finish_table(ws, AMBIGUOUS_HEADERS, AMBIGUOUS_WIDTHS)


# --- the source files, verbatim -----------------------------------------

def _write_source(ws, result: Result, side: Side) -> None:
    """The file as it arrived, with the match appended - never rewritten.

    Values come from the raw row captured at load, in the file's own column
    order, so anything the tool derived stays in the appended columns where a
    reader can tell it apart from the source.
    """
    ledger = result.gl if side is Side.GL else result.bank
    original = list(ledger.headers)
    headers = original + ["row_id", "match_id", "match_shape", "match_rule"]
    _write_header(ws, headers)

    match_by_row: dict[str, object] = {}
    for match in result.matches:
        for row_id in match.gl_ids + match.bank_ids:
            match_by_row[row_id] = match

    for excel_row, row in enumerate(ledger.rows, start=2):
        for col, name in enumerate(original, 1):
            ws.cell(row=excel_row, column=col, value=row.raw.get(name, ""))
        base = len(original)
        ws.cell(row=excel_row, column=base + 1, value=row.row_id)
        match = match_by_row.get(row.row_id)
        ws.cell(row=excel_row, column=base + 2,
                value=match.match_id if match else "")
        ws.cell(row=excel_row, column=base + 3, value=match.shape if match else "")
        ws.cell(row=excel_row, column=base + 4, value=match.rule if match else "")

    widths = [18] * len(original) + [10, 10, 13, 17]
    _finish_table(ws, headers, widths)


# --- Summary -------------------------------------------------------------

def _write_summary(ws, result: Result) -> None:
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 12
    row = 1

    def title(text: str) -> None:
        nonlocal row
        ws.cell(row=row, column=1, value=text).font = TITLE_FONT
        row += 1

    def section(text: str) -> None:
        nonlocal row
        row += 1
        ws.cell(row=row, column=1, value=text).font = SECTION_FONT
        row += 1

    def label_value(label: str, value=None) -> None:
        nonlocal row
        ws.cell(row=row, column=1, value=label)
        if value is not None:
            ws.cell(row=row, column=2, value=value)
        row += 1

    def label_money(label: str, cents: int, ruled: bool = False) -> None:
        nonlocal row
        ws.cell(row=row, column=1, value=label)
        cell = _money(ws, row, 2, cents)
        if ruled:
            cell.border = TOP_BORDER
        row += 1

    title("Reconciliation summary")
    label_value("Generated", datetime.now().strftime("%Y-%m-%d %H:%M"))

    section("FILES")
    label_value("General ledger", Path(result.gl.path).name)
    label_value("  rows", len(result.gl.rows))
    label_money("  total", result.gl.total_cents)
    label_value("Bank statement", Path(result.bank.path).name)
    label_value("  rows", len(result.bank.rows))
    label_money("  total", result.bank.total_cents)
    label_money("Difference to explain",
                result.gl.total_cents - result.bank.total_cents)

    section("MATCHED BY PASS")
    for name, count in result.pass_counts:
        label_value(name, count)
    label_value("Match groups", len(result.matches))
    label_value("GL rows matched",
                f"{result.matched_gl_count} of {len(result.gl.rows)}")
    label_value("Bank rows matched",
                f"{result.matched_bank_count} of {len(result.bank.rows)}")
    label_value("Ambiguous candidates", len(result.ambiguities))

    section("MATCH QUALITY")
    ws.cell(row=row - 1, column=2, value="groups").font = SECTION_FONT
    quality: dict[tuple[str, str, str], int] = {}
    for match in result.matches:
        key = (match.amount.value, match.date.value, match.identity.value)
        quality[key] = quality.get(key, 0) + 1
    for (amount, date_q, identity), count in sorted(quality.items()):
        label_value(f"{amount} / {date_q} / {identity}", count)

    unmatched_gl = sum(r.amount_cents for r in result.unmatched_gl)
    unmatched_bank = sum(r.amount_cents for r in result.unmatched_bank)
    open_difference = unmatched_gl - unmatched_bank
    accounted = open_difference + result.drift_cents
    difference = result.gl.total_cents - result.bank.total_cents

    section("PROOF")
    label_money(f"Open GL items ({len(result.unmatched_gl)})", unmatched_gl)
    label_money(f"Open bank items ({len(result.unmatched_bank)})", unmatched_bank)
    label_money("Open GL less open bank", open_difference)
    label_money("Drift accepted inside matches", result.drift_cents)
    label_money("Accounted for", accounted, ruled=True)
    label_money("GL total less bank total", difference)

    foots = accounted == difference
    ws.cell(row=row, column=1, value="Result")
    verdict = ws.cell(
        row=row, column=2,
        value="FOOTS - reconciliation complete" if foots
        else "DOES NOT FOOT - incomplete",
    )
    verdict.font = FOOTS_FONT if foots else FAILS_FONT
    row += 2

    ws.cell(row=row, column=1, value=(
        "Every open item plus any drift absorbed inside a match adds back to the "
        "difference the two files arrived with. That is the completeness check."
    )).alignment = Alignment(wrap_text=False)


# --- assembly ------------------------------------------------------------

def write_workbook(
    result: Result, path: str | Path, explanations: dict | None = None,
) -> Path:
    """Write the reconciliation to an .xlsx file and return where it went."""
    path = Path(path)
    wb = Workbook()

    summary = wb.active
    summary.title = "Summary"
    _write_summary(summary, result)

    _write_reconciled(wb.create_sheet("Reconciled"), result)
    _write_exceptions(wb.create_sheet("Exceptions"), result)

    # Only when there is something to show. An empty tab would imply the tool
    # had nothing to say, when in fact it had nothing to decline.
    if result.ambiguities:
        _write_ambiguous(wb.create_sheet("Ambiguous"), result, explanations)

    _write_source(wb.create_sheet("General Ledger"), result, Side.GL)
    _write_source(wb.create_sheet("Bank"), result, Side.BANK)

    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(path)
    except PermissionError:
        raise OutputLocked(
            f"Could not write {path.name} - it is probably open in Excel. "
            f"Close it and run again, or pass --out to write elsewhere."
        ) from None
    return path
